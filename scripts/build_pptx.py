"""Build the editable PowerPoint deck.

    python scripts/build_pptx.py

Produces:
    docs/presentation.pptx   — the same 4 slides as docs/KOHLER AI BathPlan Pitch Deck.pdf,
                               as native PowerPoint shapes, tables and text

The PDF deck is the fixed artefact for submission; this is the version someone
can open and edit on the day. The two are built from the same facts, and the
plan on slide 3 is drawn from the layout solver's real coordinates for the
demo brief — not a mock-up — by reusing `golden_path_layout()` from
`scripts/build_pdfs.py`.
"""

import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from build_pdfs import golden_path_layout  # noqa: E402

# ---------------------------------------------------------------------------
# Palette and type — the frontend's tokens, restated for PowerPoint
# ---------------------------------------------------------------------------

PAPER = RGBColor(0xF4, 0xF2, 0xED)
PAPER_DEEP = RGBColor(0xEB, 0xE8, 0xE1)
SURFACE = RGBColor(0xFD, 0xFC, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x19, 0x17)
INK_2 = RGBColor(0x57, 0x54, 0x4D)
INK_3 = RGBColor(0x8B, 0x88, 0x80)
RULE = RGBColor(0xDD, 0xD8, 0xCC)
ACCENT = RGBColor(0x1C, 0x4F, 0x6B)
ACCENT_SOFT = RGBColor(0xE7, 0xEF, 0xF4)
BRASS = RGBColor(0x9E, 0x75, 0x46)
OK = RGBColor(0x1F, 0x6B, 0x4A)
WARN = RGBColor(0x8A, 0x62, 0x16)

DISPLAY = "Georgia"          # headlines — editorial voice
BODY = "Segoe UI"            # prose and labels
MONO = "Consolas"            # figures and identifiers

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.72)
CONTENT_W = SLIDE_W - 2 * MARGIN

FOOTER = "KOHLER AI BathPlan — KOHLER–MIT-WPU AI Research Lab, Track 1"


# ---------------------------------------------------------------------------
# Small helpers over python-pptx
# ---------------------------------------------------------------------------

def set_fill(shape, color: RGBColor, alpha: float | None = None) -> None:
    """Solid fill, optionally translucent (python-pptx has no alpha API)."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if alpha is not None:
        solid = shape.fill._xPr.find(qn("a:solidFill"))
        srgb = solid.find(qn("a:srgbClr"))
        node = etree.SubElement(srgb, qn("a:alpha"))
        node.set("val", str(int(alpha * 100_000)))


def E(value) -> Emu:
    """Emu arithmetic produces floats; the XML writer wants whole units."""
    return Emu(int(round(value)))


def box(slide, left, top, width, height, *, fill=None, alpha=None,
        line=None, line_w=0.75, dash=None, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, E(left), E(top), E(width), E(height))
    shape.shadow.inherit = False
    if radius:
        shape.adjustments[0] = radius
    if fill is None:
        shape.fill.background()
    else:
        set_fill(shape, fill, alpha)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
        if dash:
            shape.line.dash_style = dash
    shape.text_frame.word_wrap = True
    return shape


def textbox(slide, left, top, width, height, *, anchor=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(E(left), E(top), E(width), E(height))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    return frame


def _runs(paragraph, text: str, *, size, color, font, bold, italic):
    """Render **bold** and *italic* spans as separate runs."""
    import re

    for piece in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not piece:
            continue
        run = paragraph.add_run()
        run_bold, run_italic = bold, italic
        if piece.startswith("**") and piece.endswith("**"):
            piece, run_bold = piece[2:-2], True
        elif piece.startswith("*") and piece.endswith("*"):
            piece, run_italic = piece[1:-1], True
        run.text = piece
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = font
        run.font.bold = run_bold
        run.font.italic = run_italic


def para(frame, text="", *, size=11, color=INK_2, font=BODY, bold=False,
         italic=False, space_before=0, space_after=6, line_spacing=1.28,
         align=PP_ALIGN.LEFT, indent=0.0, first=False):
    paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
    paragraph.space_before = Pt(space_before)
    paragraph.space_after = Pt(space_after)
    paragraph.line_spacing = line_spacing
    paragraph.alignment = align
    if indent:
        # _Paragraph exposes no indent API; marL + a negative first-line indent
        # is what a hanging bullet is in DrawingML.
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", str(int(Inches(indent))))
        properties.set("indent", str(-int(Inches(indent))))
    _runs(paragraph, text, size=size, color=color, font=font, bold=bold, italic=italic)
    return paragraph


def new_slide(prs, *, page: int | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_fill(box(slide, 0, 0, SLIDE_W, SLIDE_H), PAPER)
    box(slide, 0, 0, SLIDE_W, Pt(5), fill=ACCENT)
    if page is not None:
        frame = textbox(slide, MARGIN, SLIDE_H - Inches(0.44), CONTENT_W, Inches(0.24))
        para(frame, FOOTER, size=8, color=INK_3, space_after=0, first=True)
        right = textbox(slide, MARGIN, SLIDE_H - Inches(0.44), CONTENT_W, Inches(0.24))
        para(right, f"{page} / 4", size=8, color=INK_3, align=PP_ALIGN.RIGHT,
             space_after=0, first=True)
    return slide


def heading(slide, kicker: str, title: str, subtitle: str | None = None, top=Inches(0.52)):
    frame = textbox(slide, MARGIN, top, CONTENT_W, Inches(0.3))
    para(frame, kicker.upper(), size=9.5, color=BRASS, bold=True, space_after=4, first=True)
    title_frame = textbox(slide, MARGIN, top + Inches(0.3), CONTENT_W, Inches(0.6))
    para(title_frame, title, size=27, color=INK, font=DISPLAY, line_spacing=1.1,
         space_after=6, first=True)
    if subtitle:
        sub = textbox(slide, MARGIN, top + Inches(0.92), CONTENT_W, Inches(0.5))
        para(sub, subtitle, size=12, color=INK_2, line_spacing=1.35, space_after=0, first=True)
    return top + (Inches(1.5) if subtitle else Inches(1.0))


def column_rule(slide, x, top, height):
    box(slide, x, top, 9525, height, fill=RULE)


# ---------------------------------------------------------------------------
# Slide 1 — problem, solution, the numbers
# ---------------------------------------------------------------------------

STATS = [
    ("33", "curated products"),
    ("~1,460", "configurations evaluated per plan"),
    ("197", "automated tests"),
    ("0", "figures produced by an LLM"),
]


def slide_title(prs) -> None:
    slide = new_slide(prs, page=1)

    frame = textbox(slide, MARGIN, Inches(0.72), CONTENT_W, Inches(0.3))
    para(frame, "KOHLER–MIT-WPU AI RESEARCH LAB · TRACK 1", size=10, color=BRASS,
         bold=True, space_after=0, first=True)

    title = textbox(slide, MARGIN, Inches(1.06), CONTENT_W, Inches(0.9))
    para(title, "KOHLER AI BathPlan", size=42, color=INK, font=DISPLAY,
         line_spacing=1.0, space_after=0, first=True)

    sub = textbox(slide, MARGIN, Inches(1.95), Inches(9.6), Inches(0.5))
    para(sub, "The AI can imagine a bathroom. The constraint engine decides whether "
              "it survives reality.", size=14, color=ACCENT, font=DISPLAY, italic=True,
         line_spacing=1.3, space_after=0, first=True)

    box(slide, MARGIN, Inches(2.62), CONTENT_W, Pt(1), fill=RULE)

    col_w = (CONTENT_W - Inches(0.9)) / 2
    left = textbox(slide, MARGIN, Inches(2.92), col_w, Inches(2.6))
    para(left, "The problem", size=13, color=ACCENT, font=DISPLAY, bold=True,
         space_after=8, first=True)
    para(left, "People planning a bathroom do not have an imagination problem. They "
               "cannot answer the question that actually matters:", size=11, color=INK)
    para(left, "“Will these products fit in my room, together, within my budget — and "
               "if not, what should I give up?”", size=11.5, color=INK, font=DISPLAY,
         italic=True, space_before=2, space_after=8)
    para(left, "That is a configuration decision under hard constraints, not an "
               "image-generation problem. A 72-inch vanity does not fit a 5-foot wall "
               "no matter how good the render looks.", size=11, color=INK_2)

    column_rule(slide, MARGIN + col_w + Inches(0.45), Inches(2.92), Inches(2.5))

    right_x = MARGIN + col_w + Inches(0.9)
    right = textbox(slide, right_x, Inches(2.92), col_w, Inches(2.6))
    para(right, "The solution", size=13, color=ACCENT, font=DISPLAY, bold=True,
         space_after=8, first=True)
    para(right, "A deterministic planner wrapped in a conversation:", size=11, color=INK)
    for index, line in enumerate([
        "Fixtures placed against walls with the clear floor plumbing codes require",
        "Budget, compatibility, power and installation checked — unknowns stay unknown",
        "2–3 options that differ in ways you would actually choose between",
        "Water impact estimated with every formula and assumption shown",
        "**When the requirements cannot all hold, it says so — and offers the "
        "trade-offs it actually evaluated**",
    ], start=1):
        para(right, f"{index}.  {line}", size=10.5, color=INK, indent=0.24, space_after=5)

    box(slide, MARGIN, Inches(5.72), CONTENT_W, Pt(1), fill=RULE)
    stat_w = CONTENT_W / len(STATS)
    for index, (value, label) in enumerate(STATS):
        frame = textbox(slide, MARGIN + stat_w * index, Inches(5.94), stat_w - Inches(0.2),
                        Inches(0.8))
        para(frame, value, size=25, color=ACCENT, font=DISPLAY, space_after=1,
             line_spacing=1.0, first=True)
        para(frame, label, size=9, color=INK_3, space_after=0)


# ---------------------------------------------------------------------------
# Slide 2 — architecture, drawn rather than typed
# ---------------------------------------------------------------------------

def label_box(slide, left, top, width, height, lines, *, fill, line, title_color,
              title_size=10.5, body_size=8.6, radius=0.06):
    shape = box(slide, left, top, width, height, fill=fill, line=line, radius=radius)
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = Inches(0.1)
    frame.margin_top = frame.margin_bottom = Inches(0.07)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, text in enumerate(lines):
        para(frame, text, size=title_size if index == 0 else body_size,
             color=title_color if index == 0 else INK_2,
             bold=index == 0, align=PP_ALIGN.CENTER, space_after=2,
             line_spacing=1.2, first=index == 0)
    return shape


def arrow(slide, x1, y1, x2, y2, *, color=INK_3, dash=None, width=1.0):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    if dash:
        connector.line.dash_style = dash
    tail = connector.line._get_or_add_ln()
    end = etree.SubElement(tail, qn("a:tailEnd"))
    end.set("type", "triangle")
    end.set("w", "sm")
    end.set("len", "sm")
    return connector


def caption(slide, left, top, width, text, *, color=INK_3, size=8.5, align=PP_ALIGN.CENTER):
    frame = textbox(slide, left, top, width, Inches(0.26))
    para(frame, text, size=size, color=color, align=align, space_after=0, first=True)


def slide_architecture(prs) -> None:
    slide = new_slide(prs, page=2)
    heading(slide, "Architecture", "The LLM is not the source of truth",
            "Feasibility, dimensions, pricing and water figures are computed before any "
            "model is called. The model receives them as given facts.")

    top = Inches(2.08)
    # Intake row
    intake = [("Your brief + photo", MARGIN, Inches(2.6)),
              ("React + TypeScript frontend", MARGIN + Inches(2.95), Inches(3.0)),
              ("FastAPI", MARGIN + Inches(6.35), Inches(1.8))]
    for text, left, width in intake:
        label_box(slide, left, top, width, Inches(0.46), [text],
                  fill=SURFACE, line=RULE, title_color=INK, title_size=10)
    arrow(slide, MARGIN + Inches(2.6), top + Inches(0.23),
          MARGIN + Inches(2.95), top + Inches(0.23))
    arrow(slide, MARGIN + Inches(5.95), top + Inches(0.23),
          MARGIN + Inches(6.35), top + Inches(0.23))

    core_top = top + Inches(0.92)
    core_w = Inches(4.6)
    core_x = MARGIN + (CONTENT_W - core_w) / 2
    core = box(slide, core_x, core_top, core_w, Inches(1.96), fill=ACCENT_SOFT,
               line=ACCENT, line_w=1.25, radius=0.04)
    frame = core.text_frame
    frame.margin_top = Inches(0.1)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(frame, "DETERMINISTIC CORE", size=11, color=ACCENT, bold=True,
         align=PP_ALIGN.CENTER, space_after=6, first=True)
    for line in ["layout solver", "constraint engine", "recommendation engine",
                 "sustainability engine", "product catalog"]:
        para(frame, line, size=9.5, color=INK, font=MONO, align=PP_ALIGN.CENTER,
             space_after=2, line_spacing=1.15)

    side_w = Inches(2.5)
    vision = label_box(slide, MARGIN, core_top + Inches(0.26), side_w, Inches(1.44),
                       ["VISION  (Gemini)", "confidence + unknowns",
                        "authoritative: False"],
                       fill=SURFACE, line=RULE, title_color=INK_2)
    llm = label_box(slide, MARGIN + CONTENT_W - side_w, core_top + Inches(0.26), side_w,
                    Inches(1.44),
                    ["LLM LAYER  (Gemini)", "interprets + explains", "never decides"],
                    fill=SURFACE, line=RULE, title_color=INK_2)

    arrow(slide, MARGIN + side_w, core_top + Inches(0.98), core_x,
          core_top + Inches(0.98), dash=MSO_LINE_DASH_STYLE.DASH, color=WARN)
    caption(slide, MARGIN + side_w, core_top + Inches(0.62),
            core_x - (MARGIN + side_w), "advisory only", color=WARN, size=8)
    arrow(slide, core_x + core_w, core_top + Inches(0.98),
          MARGIN + CONTENT_W - side_w, core_top + Inches(0.98), color=INK_3)
    caption(slide, core_x + core_w, core_top + Inches(0.62),
            (MARGIN + CONTENT_W - side_w) - (core_x + core_w), "computed results", size=8)

    out_top = core_top + Inches(2.24)
    label_box(slide, core_x - Inches(1.1), out_top, core_w + Inches(2.2), Inches(0.44),
              ["options  ·  2D plan  ·  water impact  ·  conflicts and trade-offs"],
              fill=INK, line=INK, title_color=WHITE, title_size=10)
    arrow(slide, core_x + core_w / 2, core_top + Inches(1.96), core_x + core_w / 2, out_top)

    notes_top = Inches(5.88)
    col_w = (CONTENT_W - Inches(0.9)) / 2
    left = textbox(slide, MARGIN, notes_top, col_w, Inches(1.1))
    para(left, "Enforced structurally, not by instruction", size=11, color=ACCENT,
         font=DISPLAY, bold=True, space_after=5, first=True)
    para(left, "Vision output is typed authoritative: Literal[False] and has no field for "
               "a dimension or rough-in — it cannot reach the engine. Intent output is "
               "sanitised against a fixed vocabulary. Every response carries computed_by, "
               "and **“llm” never appears in it.**", size=9.5, color=INK_2, space_after=0)

    column_rule(slide, MARGIN + col_w + Inches(0.45), notes_top, Inches(1.0))

    right = textbox(slide, MARGIN + col_w + Inches(0.9), notes_top, col_w, Inches(1.1))
    para(right, "Works with no API key", size=11, color=ACCENT, font=DISPLAY, bold=True,
         space_after=5, first=True)
    para(right, "The AI layers are enhancements with tested fallbacks, not load-bearing "
                "parts. With no key configured, planning, the 2D plan, water estimation, "
                "conflict resolution and plain-language modification all still work — the "
                "full suite runs in this mode. A model failure degrades the wording, never "
                "the verdict.", size=9.5, color=INK_2, space_after=0)


# ---------------------------------------------------------------------------
# Slide 3 — the innovation, with the solver's real plan
# ---------------------------------------------------------------------------

CATEGORY_FILL = {
    "toilet": RGBColor(0x2F, 0x6F, 0x8F),
    "smart_toilet": RGBColor(0x1F, 0x5D, 0x7A),
    "vanity": RGBColor(0x7A, 0x5B, 0x3A),
    "basin": RGBColor(0x8A, 0x6A, 0x47),
    "shower": RGBColor(0x3F, 0x7D, 0x6A),
    "smart_shower": RGBColor(0x2F, 0x6D, 0x5A),
    "bathtub": RGBColor(0x4A, 0x6F, 0x9C),
    "storage": RGBColor(0x6B, 0x6B, 0x7D),
}

COMPARISON = [
    ["", "Typical AI tool", "BathPlan"],
    ["Spatial reasoning", "Implied by a rendering", "Solved geometry, IRC clearances"],
    ["Feasibility", "Asserted", "Computed; failing check named"],
    ["Unknowns", "Filled in plausibly", "Stay unknown, surfaced"],
    ["Impossible ask", "Something close, shown as the answer",
     "Refused, explained, alternatives"],
    ["Water savings", "A percentage", "Formula + baseline + assumptions"],
    ["LLM role", "Decides", "Interprets and explains"],
]


def draw_plan(slide, layout, left, top, width, height) -> None:
    """Draw the solver's own coordinates as native shapes — editable, not an image."""
    room_w, room_l = layout.room_width_in, layout.room_length_in
    scale = min(width / room_w, height / room_l)
    plan_w, plan_l = room_w * scale, room_l * scale
    origin_x = left + (width - plan_w) / 2
    origin_y = top + (height - plan_l)

    def to_x(value: float) -> float:
        return origin_x + value * scale

    def to_y(value: float) -> float:
        """Plan y grows away from the south wall; slides grow downward."""
        return origin_y + (room_l - value) * scale

    box(slide, origin_x, origin_y, plan_w, plan_l,
        fill=SURFACE, line=INK, line_w=1.25)

    for item in layout.placed:
        clear = item.clearance
        box(slide, to_x(clear.x_in), to_y(clear.y_in + clear.depth_in),
            clear.width_in * scale, clear.depth_in * scale,
            fill=ACCENT, alpha=0.10, line=ACCENT, line_w=0.5,
            dash=MSO_LINE_DASH_STYLE.DASH)

    if layout.door_swing:
        swing = layout.door_swing
        box(slide, to_x(swing.x_in), to_y(swing.y_in + swing.depth_in),
            swing.width_in * scale, swing.depth_in * scale,
            fill=WARN, alpha=0.12, line=WARN, line_w=0.5,
            dash=MSO_LINE_DASH_STYLE.DASH)

    for item in layout.placed:
        foot = item.footprint
        shape = box(slide, to_x(foot.x_in), to_y(foot.y_in + foot.depth_in),
                    foot.width_in * scale, foot.depth_in * scale,
                    fill=CATEGORY_FILL.get(item.category, INK_2), line=WHITE, line_w=0.75)
        if foot.width_in * scale > Inches(0.5) and foot.depth_in * scale > Inches(0.2):
            frame = shape.text_frame
            frame.margin_left = frame.margin_right = 0
            frame.margin_top = frame.margin_bottom = 0
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            para(frame, item.category.replace("_", " "), size=7, color=WHITE,
                 align=PP_ALIGN.CENTER, space_after=0, first=True)

    # Centred on the plan but free to overhang it, so it stays on one line.
    caption(slide, origin_x - Inches(0.5), origin_y + plan_l + Inches(0.06),
            plan_w + Inches(1.0),
            f"{room_w / 12:g} ft × {room_l / 12:g} ft — solver coordinates", size=8)


def slide_innovation(prs, layout, best) -> None:
    slide = new_slide(prs, page=3)
    heading(slide, "Innovation", "The most important screen is the one that says no")

    top = Inches(1.72)
    left_w = Inches(4.9)
    left = textbox(slide, MARGIN, top, left_w, Inches(2.4))
    para(left, "Constraint conflict resolution", size=12, color=ACCENT, font=DISPLAY,
         bold=True, space_after=7, first=True)
    para(left, "**User:** “Add a smart shower and a smart toilet and a bathtub, keep my "
               "budget.”", size=10.5, color=INK, space_after=5)
    para(left, "**System:** over budget by Rs 18,300 — *and* physically unplaceable in "
               "6 × 8 ft.", size=10.5, color=INK, space_after=7)
    para(left, "It names both, and notes that **more budget would not fix a spatial "
               "problem**. It then offers the relaxations it actually evaluated and lets "
               "the user choose. The plan is recomputed from scratch, not patched.",
         size=9.5, color=INK_2, space_after=6)
    para(left, "A generative tool would have returned something plausible and let the "
               "customer discover the problem during installation.", size=9.5, color=INK_2,
         space_after=0)

    table_x = MARGIN + left_w + Inches(0.5)
    table_w = MARGIN + CONTENT_W - table_x
    graphic = slide.shapes.add_table(len(COMPARISON), 3, E(table_x), E(top), E(table_w),
                                     Inches(2.3))
    table = graphic.table
    table.columns[0].width = int(table_w * 0.26)
    table.columns[1].width = int(table_w * 0.34)
    table.columns[2].width = int(table_w * 0.40)
    for row_index, row in enumerate(COMPARISON):
        table.rows[row_index].height = Inches(0.33) if row_index == 0 else Inches(0.3)
        for col_index, text in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if row_index == 0:
                cell.fill.fore_color.rgb = ACCENT
            else:
                cell.fill.fore_color.rgb = WHITE if row_index % 2 else PAPER_DEEP
            para(cell.text_frame, text, size=9,
                 color=WHITE if row_index == 0 else INK,
                 bold=row_index == 0 or col_index == 0,
                 space_after=0, line_spacing=1.1, first=True)

    plan_top = Inches(4.26)
    plan_label = textbox(slide, MARGIN, plan_top, left_w, Inches(0.3))
    para(plan_label, "Real output, not a mock-up", size=11, color=ACCENT, font=DISPLAY,
         bold=True, space_after=0, first=True)
    draw_plan(slide, layout, MARGIN, plan_top + Inches(0.3), Inches(2.05), Inches(2.05))

    notes = textbox(slide, MARGIN + Inches(2.35), plan_top + Inches(0.34),
                    left_w - Inches(2.35), Inches(2.0))
    para(notes, "The solver's actual coordinates for the demo brief. Dashed areas are the "
                "code-required clear floor and the door swing — the constraints that "
                "decided feasibility, drawn rather than asserted.", size=9.5, color=INK_2,
         space_after=6, first=True)
    para(notes, "Fixture colour is category. The room, the clearances and the swing are "
                "the solver's own numbers, redrawn as PowerPoint shapes.", size=9.5,
         color=INK_2, space_after=0)

    right_x = MARGIN + left_w + Inches(0.5)
    right = textbox(slide, right_x, plan_top + Inches(0.06), table_w, Inches(2.2))
    para(right, "Confidence-aware by construction", size=11, color=ACCENT, font=DISPLAY,
         bold=True, space_after=6, first=True)
    para(right, "A constraint report has three states, not two: *feasible*, *feasible "
                "pending verification*, and *infeasible*. An unmeasured rough-in is not "
                "the same as a mismatched one — conflating them either blocks every real "
                "plan or fakes certainty.", size=9.5, color=INK_2, space_after=6)
    para(right, "Pending options are offered, ranked strictly below verified ones, and "
                "carry their outstanding checks with them.", size=9.5, color=INK_2,
         space_after=0)

    band = box(slide, right_x, Inches(6.18), table_w, Inches(0.72), fill=ACCENT_SOFT,
               line=RULE, radius=0.05)
    frame = band.text_frame
    frame.margin_left = frame.margin_right = Inches(0.16)
    frame.margin_top = frame.margin_bottom = Inches(0.08)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(frame, f"**Demo result:** {len(layout.placed)} fixtures placed with full code "
                f"clearance, every one reachable from the doorway, at "
                f"Rs {best.total_price:,.0f} of a Rs 2,50,000 budget.",
         size=10, color=INK, space_after=0, first=True)


# ---------------------------------------------------------------------------
# Slide 4 — value to KOHLER, roadmap, and the limits stated plainly
# ---------------------------------------------------------------------------

def slide_value(prs) -> None:
    slide = new_slide(prs, page=4)
    heading(slide, "Value and roadmap", "A decision layer upstream of the design consultation")

    top = Inches(1.78)
    col_w = (CONTENT_W - Inches(0.9)) / 2

    left = textbox(slide, MARGIN, top, col_w, Inches(3.4))
    para(left, "Why this matters to KOHLER", size=12, color=ACCENT, font=DISPLAY,
         bold=True, space_after=7, first=True)
    para(left, "KOHLER already offers a human bathroom design service. BathPlan does not "
               "compete with it — it sits in front of it, turning an undecided browser "
               "into a customer arriving with a feasible, costed, explained "
               "configuration.", size=10, color=INK, space_after=8)
    for line in [
        "Specification errors surface before purchase rather than at installation.",
        "Smart and water-efficient products are recommended on evidence the customer can "
        "audit, not on marketing adjectives.",
        "Every plan produces a verification list — a natural handoff to a KOHLER "
        "consultant or installer.",
        "The catalog is designed to be swapped for an authorised live feed with no engine "
        "change.",
    ]:
        para(left, f"·   {line}", size=10, color=INK_2, indent=0.22, space_after=6)

    column_rule(slide, MARGIN + col_w + Inches(0.45), top, Inches(3.2))

    right = textbox(slide, MARGIN + col_w + Inches(0.9), top, col_w, Inches(3.4))
    para(right, "What we would build next", size=12, color=ACCENT, font=DISPLAY,
         bold=True, space_after=7, first=True)
    for index, line in enumerate([
        "**Authorised catalog integration** — replace illustrative data with a live feed; "
        "the schema already separates verified from illustrative fields.",
        "**Regional code packs** — swap IRC clearances for Indian standards. The rules are "
        "already isolated in one module.",
        "**Guided rough-in capture** — convert the commonest verification requirement into "
        "a verified fact.",
        "**Designer handoff package** — configuration, plan and open verifications, "
        "exported to a consultant.",
        "**Konnect-informed water tracking** — replace assumed usage with measured usage.",
    ], start=1):
        para(right, f"{index}.  {line}", size=10, color=INK_2, indent=0.24, space_after=7)

    band_top = Inches(5.52)
    band = box(slide, MARGIN, band_top, CONTENT_W, Inches(1.24), fill=ACCENT_SOFT,
               line=RULE, radius=0.03)
    frame = band.text_frame
    frame.margin_left = frame.margin_right = Inches(0.2)
    frame.margin_top = frame.margin_bottom = Inches(0.12)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(frame, "**Stated plainly:** prices in this prototype are illustrative and are not "
                "KOHLER pricing. Clearances are US residential-code figures used as "
                "planning minimums, and usage benchmarks are US-derived. A photograph "
                "establishes nothing measurable. This is a planning aid, not an "
                "engineering drawing — final dimensions, plumbing, electrical and "
                "local-code requirements must be verified by a qualified professional.",
         size=9.5, color=INK_2, space_after=0, first=True)


# ---------------------------------------------------------------------------

def build_pptx() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    layout, best = golden_path_layout()

    slide_title(prs)
    slide_architecture(prs)
    slide_innovation(prs, layout, best)
    slide_value(prs)

    target = REPO_ROOT / "docs" / "presentation.pptx"
    prs.save(str(target))
    return target


def main() -> int:
    deck = build_pptx()
    print(f"Wrote {deck.relative_to(REPO_ROOT)} ({deck.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
